from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "Neelesh_Mishra_Resume.docx"

NAVY = "0D1730"
NAVY_2 = "14213D"
MINT = "59E5C6"
INK = "142035"
MUTED = "5B677A"
PALE = "E9F8F4"
WHITE = "FFFFFF"
LINE = "CBD5DF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)


def set_repeat_table_layout_fixed(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_hyperlink(paragraph, text, url, color=MINT, underline=True, bold=False, size=8.2):
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_el)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(size_cs)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def style_run(run, size=8.2, color=INK, bold=False, font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold


def configure_paragraph(paragraph, before=0, after=0, line=1.0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep
    return paragraph


def add_text(cell, text, size=8.2, color=INK, bold=False, before=0, after=0, line=1.0, keep=False):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, before, after, line, keep)
    run = paragraph.add_run(text)
    style_run(run, size, color, bold)
    return paragraph


def add_section_label(cell, text, dark=False, before=6, after=4):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, before, after, 1.0, True)
    run = paragraph.add_run(text.upper())
    style_run(run, 6.9, MINT if dark else NAVY, True)
    run.font.letter_spacing = Pt(1.0) if hasattr(run.font, "letter_spacing") else None
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "7")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), MINT if dark else LINE)
    border.append(bottom)
    p_pr.append(border)
    return paragraph


def add_sidebar_item(cell, text, size=7.2, after=1.6, accent=False):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, 0, after, 1.0, True)
    run = paragraph.add_run(text)
    style_run(run, size, MINT if accent else WHITE, accent)
    return paragraph


def add_bullet(cell, text, size=7.8, color=MUTED, after=1.4, indent=0.12, hanging=0.10):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, 0, after, 1.0, True)
    paragraph.paragraph_format.left_indent = Inches(indent)
    paragraph.paragraph_format.first_line_indent = Inches(-hanging)
    lead = paragraph.add_run("- ")
    style_run(lead, size, MINT, True)
    run = paragraph.add_run(text)
    style_run(run, size, color)
    return paragraph


def add_job_header(cell, role, company, period, location, before=3):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, before, 0.5, 1.0, True)
    role_run = paragraph.add_run(role)
    style_run(role_run, 10.3, NAVY, True)
    company_run = paragraph.add_run(f"  |  {company}")
    style_run(company_run, 8.1, MINT, True)
    meta = add_text(cell, f"{period}  |  {location}", 7.2, MUTED, False, 0, 3, 1.0, True)
    return paragraph, meta


def add_project(cell, title, description, stack, before=1.2):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, before, 0.4, 1.0, True)
    mark = paragraph.add_run("◆ ")
    style_run(mark, 6.2, MINT, True)
    title_run = paragraph.add_run(title)
    style_run(title_run, 8.0, NAVY, True)
    desc = add_text(cell, description, 7.25, INK, False, 0, 0.5, 1.0, True)
    desc.paragraph_format.left_indent = Inches(0.12)
    stack_p = add_text(cell, stack, 6.55, MUTED, False, 0, 1.3, 1.0, True)
    stack_p.paragraph_format.left_indent = Inches(0.12)


def add_ai_initiative(cell, title, subtitle, url, bullets, stack):
    paragraph = cell.add_paragraph()
    configure_paragraph(paragraph, 3, 0.4, 1.0, True)
    add_hyperlink(paragraph, title, url, color=NAVY, underline=False, bold=True, size=10.1)
    sub_run = paragraph.add_run(f"  |  {subtitle}  |  2026")
    style_run(sub_run, 7.5, MINT, True)
    for bullet in bullets:
        add_bullet(cell, bullet, 7.65, INK, 1.6, 0.16, 0.10)
    stack_p = add_text(cell, stack, 6.75, MUTED, False, 0, 3, 1.0, True)
    stack_p.paragraph_format.left_indent = Inches(0.16)


def remove_empty_first_paragraph(cell):
    paragraphs = cell.paragraphs
    if paragraphs and not paragraphs[0].text and len(paragraphs) > 1:
        p = paragraphs[0]._element
        p.getparent().remove(p)


def add_page(document, page_number):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_no_borders(table)
    set_repeat_table_layout_fixed(table)
    table.columns[0].width = Inches(1.95)
    table.columns[1].width = Inches(5.95)
    row = table.rows[0]
    row.height = Inches(10.4)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    sidebar, main = row.cells
    sidebar.width = Inches(1.95)
    main.width = Inches(5.95)
    sidebar.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_fill(sidebar, NAVY)
    set_cell_fill(main, WHITE)
    set_cell_margins(sidebar, 220, 210, 180, 210)
    set_cell_margins(main, 220, 260, 180, 250)
    if page_number == 1:
        build_page_one(sidebar, main)
    else:
        build_page_two(sidebar, main)
    remove_empty_first_paragraph(sidebar)
    remove_empty_first_paragraph(main)


def build_page_one(sidebar, main):
    p = sidebar.add_paragraph()
    configure_paragraph(p, 0, 1, 0.88, True)
    r = p.add_run("NEELESH\nMISHRA")
    style_run(r, 20, WHITE, True)
    add_sidebar_item(sidebar, "SENIOR SOFTWARE ENGINEER", 6.8, 9, True)

    add_section_label(sidebar, "Contact", True, 0, 4)
    for line in ("Dubai, UAE", "+971-502416958", "+91-9999695408", "n.mish2202@gmail.com"):
        add_sidebar_item(sidebar, line)
    p = sidebar.add_paragraph()
    configure_paragraph(p, 0, 2, 1.0, True)
    add_hyperlink(p, "LinkedIn profile", "https://linkedin.com/in/neelesh-mishra-6b5066108", size=7.2)
    p = sidebar.add_paragraph()
    configure_paragraph(p, 0, 2, 1.0, True)
    add_hyperlink(p, "GitHub profile", "https://github.com/nmish2202", size=7.2)

    add_section_label(sidebar, "Core stack", True)
    for item in ("React / Next.js", "Node.js / TypeScript", "PHP / Laravel / CodeIgniter", "Tailwind / MUI / shadcn/ui", "PostgreSQL / MySQL / MongoDB"):
        add_sidebar_item(sidebar, item)

    add_section_label(sidebar, "Cloud & delivery", True)
    for item in ("AWS (Certified)", "Docker", "CI/CD pipelines", "Prisma ORM"):
        add_sidebar_item(sidebar, item)

    add_section_label(sidebar, "Applied AI", True)
    for item in ("Retrieval-Augmented Generation", "Azure OpenAI", "Vector search / pgvector", "Local embeddings / FastEmbed", "RBAC-aware retrieval"):
        add_sidebar_item(sidebar, item)

    add_section_label(sidebar, "Certifications", True)
    for item in ("AWS Certified", "Davra IoT Platform Certification"):
        add_sidebar_item(sidebar, item)

    add_section_label(sidebar, "Education", True)
    add_sidebar_item(sidebar, "MCA", accent=True)
    add_sidebar_item(sidebar, "Mangalayatan University, 2024", 6.8)
    add_sidebar_item(sidebar, "B.Sc. Mathematics", accent=True)
    add_sidebar_item(sidebar, "CSJM University, 2014-2017", 6.8)

    add_section_label(sidebar, "Languages", True)
    add_sidebar_item(sidebar, "Hindi - Full Professional", 6.9)
    add_sidebar_item(sidebar, "English - Professional Working", 6.9)

    add_section_label(main, "Profile", False, 0, 4)
    add_text(
        main,
        "Senior full stack developer with 8+ years building government and enterprise platforms using React, Next.js, Node.js, PHP, and cloud delivery practices. Hands-on with citation-grounded RAG, local embeddings, pgvector retrieval, Azure OpenAI, and authorization-aware AI systems. Currently delivering secure, multilingual products at e& enterprise in Dubai.",
        8.05, INK, False, 0, 5, 1.03, True,
    )
    add_section_label(main, "Experience", False, 2, 3)
    add_job_header(main, "Senior Software Engineer", "e& enterprise", "November 2022 - Present", "Dubai, UAE", 1)

    add_project(main, "Oyoon - Dubai Police Camera Monitoring", "Built the administration platform controlling officer access to camera feeds and locations, with user and group management, granular permissions, Active Directory SSO, ESB notifications, and complete audit history.", "Next.js | TypeScript | Tailwind CSS | Active Directory | ESB")
    add_project(main, "SmartKhateeb - AWQAF", "Built trilingual scheduling, evaluation, recording, and reporting workflows supporting 1,000+ khateebs across all seven UAE Emirates.", "Next.js 15 | shadcn/ui | TanStack Table | Zustand | RTL/i18n")
    add_project(main, "Sharjah Safari Website", "Developed a bilingual custom WordPress experience with 15+ ACF Pro content blocks, WPML localization, and an interactive park map.", "WordPress | PHP | ACF Pro | WPML | MySQL")
    add_project(main, "MPOS - Abu Dhabi Commercial Bank", "Delivered dual portals for merchant requests, POS inventory, SLA visibility, and field-team assignment workflows.", "Next.js 14 | MUI | React Hook Form | Redux | Framer Motion")
    add_project(main, "ADNOC NOC Portal", "Shipped a permitting portal for smart infrastructure across 72 ADNOC sites in all seven Emirates, including a five-stage approval chain and Mapbox installation view.", "Next.js | MUI | Mapbox | ApexCharts | RTL/i18n")
    add_project(main, "NHRI Website & Complaint Portal", "Built the bilingual public website, CMS administration, and structured citizen complaint experience for the UAE National Human Rights Institution.", "Laravel | PHP | Tailwind CSS | MySQL | RTL/i18n")
    add_project(main, "Oyoon SIM & Router Management", "Delivered a four-role hardware logistics platform with two-stage approvals, live availability, overdue alerts, and issuance and return tracking.", "Next.js 15 | TypeScript | Prisma | PostgreSQL | shadcn/ui | Docker")
    add_project(main, "GCGRA Self-Exclusion Portal", "Built an accessible national self-exclusion journey with UAE Pass authentication, five gaming categories, six-month to five-year periods, and WCAG 2.1 AA-aligned interfaces.", "Next.js 15 | Tailwind CSS | Radix UI | UAE Pass | WCAG")


def build_page_two(sidebar, main):
    add_sidebar_item(sidebar, "02 / 02", 7.1, 10, True)
    add_section_label(sidebar, "AI toolkit", True, 0, 5)
    for item in ("Document ingestion", "Metadata-aware chunking", "BGE / sentence-transformers", "HNSW vector retrieval", "Grounded generation", "Page & section citations", "JWT authentication", "Database-enforced RBAC", "Streaming responses", "Role-tagged batch ingestion"):
        add_sidebar_item(sidebar, item, 7.1, 2.0)

    add_section_label(sidebar, "Architecture", True, 8, 5)
    for item in ("Next.js", "FastAPI", "PostgreSQL / pgvector", "Azure OpenAI", "Docker Compose", "Alembic migrations", "pytest / Jest"):
        add_sidebar_item(sidebar, item, 7.1, 2.0)

    add_section_label(sidebar, "Principles", True, 8, 5)
    for item in ("Citations over unsupported claims", "Permissions before model context", "Local embeddings where practical", "Prototype scope stated clearly"):
        add_sidebar_item(sidebar, item, 7.0, 2.5)

    add_section_label(sidebar, "Availability", True, 8, 5)
    add_sidebar_item(sidebar, "UAE, India, and international remote roles", 7.2, 3, True)
    add_sidebar_item(sidebar, "neeleshmishra.dev", 7.0)

    add_section_label(main, "AI Initiatives", False, 0, 3)
    add_text(main, "Personal, publicly shareable prototypes. Project names link to the source repositories.", 7.3, MUTED, False, 0, 3, 1.0, True)
    add_ai_initiative(
        main,
        "Enterprise RAG",
        "Document-Grounded Q&A",
        "https://github.com/nmish2202/Enterprise-RAG",
        [
            "Built a full-stack RAG prototype supporting PDF, DOCX, and TXT ingestion, metadata-aware chunking, local BGE embeddings, pgvector HNSW retrieval, and Azure OpenAI answers with page and section citations.",
            "Containerized the Next.js, FastAPI, and PostgreSQL stack with automatic migrations and backend and frontend test suites.",
        ],
        "Next.js | FastAPI | PostgreSQL | pgvector | FastEmbed | Azure OpenAI | Docker",
    )
    add_ai_initiative(
        main,
        "Enterprise RBAC RAG",
        "Authorization-Aware Retrieval",
        "https://github.com/nmish2202/Enterprise-RBAC-RAG",
        [
            "Extended the RAG architecture with JWT authentication and database-enforced RBAC inside vector search, preventing unauthorized document content from entering retrieved context or citations.",
            "Added streamed answers, role-tagged batch ingestion, immediate account and role revocation, and user and role administration using Next.js 15, asynchronous FastAPI, PostgreSQL/pgvector, local sentence-transformers, and Azure OpenAI.",
        ],
        "Next.js 15 | Async FastAPI | PostgreSQL | pgvector | Sentence Transformers | Azure OpenAI",
    )

    add_section_label(main, "Earlier Experience", False, 5, 3)
    add_job_header(main, "Software Developer", "Appventurez", "June 2021 - November 2022", "Noida, India", 1)
    for item in (
        "Developed full-stack web and mobile-backend solutions for client projects using React.js, Node.js, and modern JavaScript frameworks.",
        "Built and integrated REST APIs connecting front-end applications with backend services and third-party platforms.",
        "Worked directly with clients and project managers to translate business requirements into technical specifications.",
    ):
        add_bullet(main, item, 7.7, INK, 1.6)

    add_job_header(main, "PHP Developer", "TechGropse Pvt. Ltd.", "October 2017 - May 2021", "Noida, India", 6)
    for item in (
        "Developed and maintained dynamic web applications using PHP frameworks including Laravel and CodeIgniter.",
        "Designed and optimized MySQL database schemas and queries to support application features and reporting.",
        "Built responsive front-end interfaces for cross-browser compatibility, growing into increasing ownership of features over more than three years.",
    ):
        add_bullet(main, item, 7.7, INK, 1.6)

    add_section_label(main, "Delivery Profile", False, 8, 3)
    add_text(main, "Secure enterprise workflows | multilingual and RTL interfaces | accessible product systems | cloud and container delivery | frontend-to-backend ownership", 7.5, INK, False, 0, 2, 1.0, True)

    add_section_label(main, "AI Architecture Evidence", False, 8, 3)
    add_project(main, "Grounded retrieval path", "Document parsing and page or section metadata flow into local embeddings and pgvector HNSW search; Azure OpenAI receives only retrieved context and returns answers with source citations.", "Ingestion -> chunk metadata -> local vectors -> HNSW retrieval -> grounded answer -> citations", 1)
    add_project(main, "Authorization boundary", "JWT identity and current database role grants constrain the vector-ranking query itself, so unauthorized content cannot enter model context or citation output.", "Identity -> active role lookup -> permitted documents -> ranked context -> streamed answer", 1)
    add_project(main, "Prototype operations", "Docker Compose, automatic migrations, test suites, content-hash refresh, and immediate account and role checks make both initiatives repeatable and inspectable.", "Containers | migrations | backend and frontend tests | revocation checks", 1)


def build_document():
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.3)
    section.bottom_margin = Inches(0.3)
    section.left_margin = Inches(0.3)
    section.right_margin = Inches(0.3)
    section.header_distance = Inches(0.1)
    section.footer_distance = Inches(0.1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(0)

    props = document.core_properties
    props.title = "Neelesh Mishra - Senior Software Engineer Resume"
    props.subject = "Full stack engineering, enterprise delivery, and AI initiatives"
    props.author = "Neelesh Mishra"
    props.keywords = "Next.js, Full Stack, RAG, Azure OpenAI, pgvector, RBAC, Dubai"
    props.comments = "Two-page professional resume with verified AI initiative links."

    add_page(document, 1)
    add_page(document, 2)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
